from __future__ import annotations

import json
from datetime import UTC, datetime
from io import BytesIO
from statistics import quantiles
from time import perf_counter

from openpyxl import load_workbook

from app.agent.dependencies import build_dependencies
from app.agent.graph import build_main_graph
from app.delivery.models import DeliveryBundle
from app.exports.registry import ExportFieldRegistry
from app.exports.workbook import build_workbook


def expand_bundle(bundle: DeliveryBundle, count: int) -> DeliveryBundle:
    leads = []
    details = {}
    for index in range(count):
        source = bundle.leads[index % len(bundle.leads)]
        enterprise_id = f"{source.enterprise_id}-{index}"
        lead = source.model_copy(
            deep=True,
            update={
                "rank": index + 1,
                "enterprise_id": enterprise_id,
                "enterprise_name": f"{source.enterprise_name}-{index + 1}",
            },
        )
        source_detail = bundle.details[source.enterprise_id]
        details[enterprise_id] = source_detail.model_copy(
            deep=True,
            update={
                "lead": lead,
                "score": source_detail.score.model_copy(
                    deep=True, update={"enterprise_id": enterprise_id}
                ),
            },
        )
        leads.append(lead)
    return bundle.model_copy(
        deep=True,
        update={
            "snapshot": bundle.snapshot.model_copy(update={"result_count": count}),
            "task": bundle.task.model_copy(update={"result_count": count}),
            "leads": leads,
            "details": details,
        },
    )


def completed_bundle():
    deps = build_dependencies()
    graph = build_main_graph(deps)
    session_id = "phase7-performance"
    graph.invoke(
        {
            "session_id": session_id,
            "incoming_text": "帮我找上海松江3家集团V网企业",
        },
        config={"configurable": {"thread_id": session_id}},
    )
    task = deps.task_repository.get_active_task(session_id)
    bundle = deps.delivery_query_service.freeze(task.task_id, task.version)
    return deps, task, bundle


def benchmark_workbook(bundle: DeliveryBundle, count: int) -> dict[str, float | int]:
    expanded = expand_bundle(bundle, count)
    fields = ExportFieldRegistry().validate(
        ["rank", "enterprise_name", "industry", "lead_score", "verification_status"]
    )
    started = perf_counter()
    content = build_workbook(
        expanded,
        fields,
        target_count=count,
        export_id=f"benchmark-{count}",
        exported_at=datetime.now(UTC),
        include_task_summary=True,
        include_score_breakdown=True,
        include_evidence_summary=True,
        include_conflicts=True,
    )
    elapsed_ms = (perf_counter() - started) * 1000
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)
    return {
        "rows": count,
        "generation_ms": round(elapsed_ms, 3),
        "size_bytes": len(content),
        "sheet_count": len(workbook.sheetnames),
        "lead_sheet_rows": workbook["潜客清单"].max_row - 1,
        "evidence_sheet_rows": workbook["证据摘要"].max_row - 1,
    }


def run_benchmark() -> dict:
    deps, task, bundle = completed_bundle()
    expanded = expand_bundle(bundle, 1000)
    deps.delivery_snapshot_repository.save(expanded)
    samples = []
    for index in range(100):
        started = perf_counter()
        deps.delivery_query_service.lead_page(
            task.task_id,
            task.version,
            page=index % 10 + 1,
            page_size=20,
            sort_by="lead_score" if index % 2 else "rank",
            sort_order="desc" if index % 2 else "asc",
        )
        samples.append((perf_counter() - started) * 1000)
    return {
        "measured_at": datetime.now(UTC).isoformat(),
        "environment": "local in-memory delivery repository",
        "lead_list": {
            "dataset_rows": 1000,
            "samples": len(samples),
            "p95_ms": round(quantiles(samples, n=100)[94], 3),
        },
        "workbook_100": benchmark_workbook(bundle, 100),
        "workbook_1000": benchmark_workbook(bundle, 1000),
    }


if __name__ == "__main__":
    print(json.dumps(run_benchmark(), ensure_ascii=False, indent=2))

