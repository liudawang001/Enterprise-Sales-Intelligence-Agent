from __future__ import annotations

from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.delivery.models import DeliveryBundle
from app.exports.registry import ExportFieldDefinition
from app.exports.sanitizer import safe_http_url, safe_spreadsheet_text

HEADER_FILL = PatternFill("solid", fgColor="1F4E78")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUBHEADER_FILL = PatternFill("solid", fgColor="D9EAF7")
STATUS_FILLS = {
    "VERIFIED": PatternFill("solid", fgColor="E2F0D9"),
    "PARTIAL": PatternFill("solid", fgColor="FFF2CC"),
    "CONFLICTING": PatternFill("solid", fgColor="FCE4D6"),
    "UNVERIFIED": PatternFill("solid", fgColor="E7E6E6"),
    "MISSING": PatternFill("solid", fgColor="E7E6E6"),
}


def _excel_value(value: Any) -> Any:
    if value is None:
        return "MISSING"
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (int, float, bool)):
        return value
    return safe_spreadsheet_text(value)


def _write_cell(cell: Cell, value: Any, *, data_type: str = "TEXT") -> None:
    cell.value = _excel_value(value)
    cell.alignment = Alignment(vertical="top", wrap_text=True)
    if data_type == "PERCENT" and isinstance(cell.value, (int, float)):
        cell.number_format = "0.0%"
    elif data_type == "NUMBER" and isinstance(cell.value, (int, float)):
        cell.number_format = "0.00"
    elif data_type == "INTEGER" and isinstance(cell.value, int):
        cell.number_format = "#,##0"
    elif data_type == "DATETIME" and isinstance(cell.value, datetime):
        cell.number_format = "yyyy-mm-dd hh:mm"
    if data_type == "URL":
        url = safe_http_url(value)
        if url:
            cell.value = safe_spreadsheet_text(url)
            cell.hyperlink = url
            cell.style = "Hyperlink"


def _style_table(sheet, width_hints: dict[int, int] | None = None) -> None:
    if sheet.max_row >= 1:
        for cell in sheet[1]:
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = Alignment(vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        sheet.row_dimensions[1].height = 30
    for column in range(1, sheet.max_column + 1):
        values = [
            len(str(sheet.cell(row=row, column=column).value or ""))
            for row in range(1, min(sheet.max_row, 60) + 1)
        ]
        width = min(max(max(values, default=10) + 2, 10), 36)
        if width_hints and column in width_hints:
            width = width_hints[column]
        sheet.column_dimensions[get_column_letter(column)].width = width


def _lead_sheet(workbook: Workbook, bundle: DeliveryBundle, fields: list[ExportFieldDefinition], count: int) -> None:
    sheet = workbook.active
    sheet.title = "潜客清单"
    for column, field in enumerate(fields, start=1):
        sheet.cell(row=1, column=column, value=field.display_name)
    for row_index, lead in enumerate(bundle.leads[:count], start=2):
        for column, field in enumerate(fields, start=1):
            _write_cell(
                sheet.cell(row=row_index, column=column),
                getattr(lead, field.source_path),
                data_type=field.data_type,
            )
            if field.field_key == "verification_status":
                fill = STATUS_FILLS.get(lead.verification_status)
                if fill:
                    sheet.cell(row=row_index, column=column).fill = fill
    _style_table(sheet)


def _score_sheet(workbook: Workbook, bundle: DeliveryBundle, count: int) -> None:
    sheet = workbook.create_sheet("评分说明")
    components = []
    for detail in list(bundle.details.values())[:count]:
        for item in detail.score.components:
            if item["component"] not in components:
                components.append(item["component"])
    headers = ["企业名称", "总分", *components, "Scoring Profile Version"]
    sheet.append(headers)
    for lead in bundle.leads[:count]:
        detail = bundle.details[lead.enterprise_id]
        scores = {
            item["component"]: item["weighted_score"]
            for item in detail.score.components
        }
        sheet.append(
            [
                safe_spreadsheet_text(lead.enterprise_name),
                lead.lead_score,
                *[scores.get(name, "MISSING") for name in components],
                detail.score.scoring_profile_version,
            ]
        )
    for row in sheet.iter_rows(min_row=2, min_col=2, max_col=2 + len(components)):
        for cell in row:
            if isinstance(cell.value, (int, float)):
                cell.number_format = "0.00"
    _style_table(sheet)


def _evidence_sheet(workbook: Workbook, bundle: DeliveryBundle, count: int) -> None:
    sheet = workbook.create_sheet("证据摘要")
    headers = [
        "Enterprise",
        "Field",
        "Primary Value",
        "Verification Status",
        "Confidence",
        "Primary Source Type",
        "Provider",
        "Source URL",
        "Retrieved At",
        "Supporting Evidence Count",
        "Conflict Count",
    ]
    sheet.append(headers)
    row_index = 2
    for lead in bundle.leads[:count]:
        detail = bundle.details[lead.enterprise_id]
        evidence_by_id = {item.evidence_id: item for item in detail.evidence}
        for field in detail.resolved_fields:
            supporting = field.get("supporting_evidence_ids", [])
            conflicting = field.get("conflicting_evidence_ids", [])
            primary = next(
                (evidence_by_id[item] for item in supporting if item in evidence_by_id),
                None,
            )
            values = [
                lead.enterprise_name,
                field["field_name"],
                field.get("primary_value"),
                field["status"],
                field["confidence"],
                primary.source_type if primary else None,
                primary.provider if primary else None,
                primary.source_url if primary else None,
                primary.retrieved_at if primary else None,
                len(supporting),
                len(conflicting),
            ]
            for column, value in enumerate(values, start=1):
                kind = "URL" if column == 8 else "DATETIME" if column == 9 else "PERCENT" if column == 5 else "TEXT"
                _write_cell(sheet.cell(row_index, column), value, data_type=kind)
            row_index += 1
    _style_table(sheet)


def _task_sheet(workbook: Workbook, bundle: DeliveryBundle, fields: list[ExportFieldDefinition], export_id: str, exported_at: datetime) -> None:
    sheet = workbook.create_sheet("任务信息")
    sheet.append(["项目", "值"])
    task = bundle.task
    values = [
        ("Project Name", "Enterprise Sales Intelligence Agent"),
        ("Generated by", "Enterprise Sales Intelligence Agent"),
        ("Data Scope", "Public enterprise data + configured data sources"),
        ("task_id", task.task_id),
        ("task_version", task.viewed_version),
        ("business", task.business),
        ("region", task.region),
        ("target_count", task.target_count),
        ("criteria_snapshot_id", task.criteria_snapshot_id),
        ("verified_lead_set_id", task.verified_lead_set_id),
        ("lead_score_set_id", task.lead_score_set_id),
        ("scoring_profile_id", bundle.snapshot.scoring_profile_id),
        ("delivery_snapshot_id", bundle.snapshot.snapshot_id),
        ("export_id", export_id),
        ("export_time", exported_at),
        ("export_fields", ", ".join(item.field_key for item in fields)),
        ("Hard Constraints", str(task.hard_constraints)),
        ("Soft Preferences", str(task.soft_preferences)),
    ]
    for label, value in values:
        sheet.append([label, _excel_value(value)])
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 72
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for row in sheet.iter_rows(min_row=2):
        row[0].font = Font(bold=True)
        row[0].fill = SUBHEADER_FILL
        row[1].alignment = Alignment(vertical="top", wrap_text=True)


def _conflict_sheet(workbook: Workbook, bundle: DeliveryBundle, count: int) -> None:
    sheet = workbook.create_sheet("冲突明细")
    headers = [
        "Enterprise",
        "Field",
        "Primary Value",
        "Alternative Value",
        "Source Type",
        "Provider",
        "Source URL",
        "Confidence",
        "Retrieved At",
    ]
    sheet.append(headers)
    for lead in bundle.leads[:count]:
        detail = bundle.details[lead.enterprise_id]
        evidence_by_id = {item.evidence_id: item for item in detail.evidence}
        for field in detail.resolved_fields:
            if field["status"] != "CONFLICTING":
                continue
            alternatives = field.get("alternatives") or ["MISSING"]
            conflicts = field.get("conflicting_evidence_ids") or [None]
            for index, alternative in enumerate(alternatives):
                evidence = evidence_by_id.get(conflicts[min(index, len(conflicts) - 1)])
                row = [
                    lead.enterprise_name,
                    field["field_name"],
                    field.get("primary_value"),
                    alternative,
                    evidence.source_type if evidence else None,
                    evidence.provider if evidence else None,
                    evidence.source_url if evidence else None,
                    evidence.confidence if evidence else field["confidence"],
                    evidence.retrieved_at if evidence else field.get("resolved_at"),
                ]
                sheet.append([_excel_value(value) for value in row])
                url = safe_http_url(row[6])
                if url:
                    sheet.cell(sheet.max_row, 7).hyperlink = url
                    sheet.cell(sheet.max_row, 7).style = "Hyperlink"
                sheet.cell(sheet.max_row, 8).number_format = "0.0%"
                if isinstance(sheet.cell(sheet.max_row, 9).value, datetime):
                    sheet.cell(sheet.max_row, 9).number_format = "yyyy-mm-dd hh:mm"
    _style_table(sheet)


def build_workbook(
    bundle: DeliveryBundle,
    fields: list[ExportFieldDefinition],
    *,
    target_count: int,
    export_id: str,
    exported_at: datetime,
    include_task_summary: bool,
    include_score_breakdown: bool,
    include_evidence_summary: bool,
    include_conflicts: bool,
) -> bytes:
    workbook = Workbook()
    _lead_sheet(workbook, bundle, fields, target_count)
    if include_score_breakdown:
        _score_sheet(workbook, bundle, target_count)
    if include_evidence_summary:
        _evidence_sheet(workbook, bundle, target_count)
    if include_task_summary:
        _task_sheet(workbook, bundle, fields, export_id, exported_at)
    if include_conflicts:
        _conflict_sheet(workbook, bundle, target_count)
    workbook.calculation.fullCalcOnLoad = True
    stream = BytesIO()
    workbook.save(stream)
    return stream.getvalue()
