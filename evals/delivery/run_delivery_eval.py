from __future__ import annotations

import json
from dataclasses import asdict, dataclass
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import load_workbook

from app.agent.dependencies import build_dependencies
from app.agent.graph import build_main_graph
from app.domain.task import TaskPatch
from app.exports.models import CreateExportRequest
from app.exports.registry import ExportFieldRegistry
from app.exports.sanitizer import (
    safe_http_url,
    safe_spreadsheet_text,
    sanitize_filename,
)
from app.exports.workbook import build_workbook


@dataclass
class DeliveryEvalReport:
    snapshot_case_count: int
    export_field_case_count: int
    evidence_case_count: int
    historical_case_count: int
    security_case_count: int
    snapshot_version_accuracy: float
    ui_excel_critical_field_consistency: float
    export_field_accuracy: float
    export_field_allowlist_enforcement: float
    evidence_traceability: float
    export_idempotency: float
    workbook_parse_success: float
    unsafe_hyperlink_rejection: float
    formula_injection_protection: float
    filename_traversal_protection: float
    historical_version_accuracy: float


def _completed_bundle(session_id: str):
    deps = build_dependencies()
    graph = build_main_graph(deps)
    graph.invoke(
        {
            "session_id": session_id,
            "incoming_text": "帮我找上海松江3家集团V网企业",
        },
        config={"configurable": {"thread_id": session_id}},
    )
    task = deps.task_repository.get_active_task(session_id)
    return deps, task, deps.delivery_query_service.freeze(task.task_id, task.version)


def _cell_value(value: Any) -> Any:
    if value is None:
        return "MISSING"
    if isinstance(value, datetime):
        return value.replace(tzinfo=None) if value.tzinfo else value
    if isinstance(value, (int, float, bool)):
        return value
    return safe_spreadsheet_text(value)


def _cell_matches(actual: Any, expected: Any) -> bool:
    if isinstance(actual, datetime) and isinstance(expected, datetime):
        return abs((actual - expected).total_seconds()) < 0.001
    return actual == expected


def _snapshot_cases(deps, bundle) -> tuple[int, int]:
    lead = bundle.leads[0]
    detail = bundle.details[lead.enterprise_id]
    before = [
        bundle.snapshot.snapshot_id,
        bundle.snapshot.task_id,
        bundle.snapshot.task_version,
        bundle.snapshot.verified_lead_set_id,
        bundle.snapshot.lead_score_set_id,
        bundle.snapshot.result_count,
        bundle.task.business,
        bundle.task.region,
        bundle.task.target_count,
        lead.rank,
        lead.enterprise_id,
        lead.enterprise_name,
        lead.industry,
        lead.region,
        lead.office_count,
        lead.public_phone,
        lead.website,
        lead.lead_score,
        lead.verification_status,
        tuple(item.evidence_id for item in detail.evidence),
    ]
    enterprise = deps.enterprise_repository.enterprises[lead.enterprise_id]
    deps.enterprise_repository.enterprises[lead.enterprise_id] = enterprise.model_copy(
        update={"canonical_name": "快照冻结后变更的底层名称"}
    )
    reread = deps.delivery_query_service.get_bundle(bundle.snapshot.snapshot_id)
    reread_lead = reread.leads[0]
    reread_detail = reread.details[reread_lead.enterprise_id]
    after = [
        reread.snapshot.snapshot_id,
        reread.snapshot.task_id,
        reread.snapshot.task_version,
        reread.snapshot.verified_lead_set_id,
        reread.snapshot.lead_score_set_id,
        reread.snapshot.result_count,
        reread.task.business,
        reread.task.region,
        reread.task.target_count,
        reread_lead.rank,
        reread_lead.enterprise_id,
        reread_lead.enterprise_name,
        reread_lead.industry,
        reread_lead.region,
        reread_lead.office_count,
        reread_lead.public_phone,
        reread_lead.website,
        reread_lead.lead_score,
        reread_lead.verification_status,
        tuple(item.evidence_id for item in reread_detail.evidence),
    ]
    return len(before), sum(left == right for left, right in zip(before, after))


def _export_field_cases(bundle) -> tuple[int, int]:
    registry = ExportFieldRegistry()
    public = registry.public_fields()
    cases = [[item.field_key] for item in public]
    cases.extend(
        [
            ["rank", "enterprise_name", "lead_score", "verification_status"],
            ["public_phone", "website", "primary_source_url"],
        ]
    )
    correct = 0
    for index, keys in enumerate(cases):
        fields = registry.validate(keys)
        content = build_workbook(
            bundle,
            fields,
            target_count=1,
            export_id=f"field-case-{index}",
            exported_at=bundle.snapshot.created_at,
            include_task_summary=False,
            include_score_breakdown=False,
            include_evidence_summary=False,
            include_conflicts=False,
        )
        workbook = load_workbook(BytesIO(content), data_only=False)
        sheet = workbook["潜客清单"]
        headers = [sheet.cell(1, column).value for column in range(1, len(fields) + 1)]
        values = [sheet.cell(2, column).value for column in range(1, len(fields) + 1)]
        expected = [_cell_value(getattr(bundle.leads[0], item.source_path)) for item in fields]
        correct += headers == [item.display_name for item in fields] and all(
            _cell_matches(actual, wanted)
            for actual, wanted in zip(values, expected)
        )
    return len(cases), correct


def _evidence_cases(bundle) -> tuple[int, int]:
    cases = []
    for detail in bundle.details.values():
        resolved = {item["field_name"]: item for item in detail.resolved_fields}
        for evidence in detail.evidence:
            cases.append((detail, resolved, evidence))
    cases = cases[:20]
    correct = 0
    for detail, resolved, evidence in cases:
        field = resolved.get(evidence.field_name)
        linked_ids = set(field.get("supporting_evidence_ids", [])) | set(
            field.get("conflicting_evidence_ids", [])
        ) if field else set()
        correct += bool(
            field
            and evidence.evidence_id in linked_ids
            and evidence.source_record_id
            and evidence.provider
            and evidence.source_type
            and evidence.retrieved_at
            and (
                evidence.source_url is None
                or safe_http_url(evidence.source_url) == evidence.source_url
            )
            and detail.lead.field_statuses.get(evidence.field_name) == field["status"]
        )
    return len(cases), correct


def _historical_cases(deps, task, bundle) -> tuple[int, int]:
    old_snapshot = bundle.snapshot
    old_leads = bundle.leads
    deps.task_service.apply_patch(
        task.task_id,
        TaskPatch(target_count=max(1, (task.target_count or 2) - 1)),
    )
    current = deps.task_repository.get_task(task.task_id)
    historical = deps.delivery_query_service.freeze(task.task_id, task.version)
    checks = [
        current.version == task.version + 1,
        historical.snapshot.snapshot_id == old_snapshot.snapshot_id,
        historical.snapshot.task_version == task.version,
        historical.snapshot.verified_lead_set_id == old_snapshot.verified_lead_set_id,
        historical.snapshot.lead_score_set_id == old_snapshot.lead_score_set_id,
        historical.snapshot.result_count == old_snapshot.result_count,
        [item.rank for item in historical.leads] == [item.rank for item in old_leads],
        [item.lead_score for item in historical.leads]
        == [item.lead_score for item in old_leads],
        [item.verification_status for item in historical.leads]
        == [item.verification_status for item in old_leads],
        [
            item.evidence_id
            for detail in historical.details.values()
            for item in detail.evidence
        ]
        == [
            item.evidence_id
            for detail in bundle.details.values()
            for item in detail.evidence
        ],
    ]
    return len(checks), sum(checks)


def _security_cases() -> tuple[int, int, int, int]:
    formulas = ('=HYPERLINK("https://evil.example")', "+1+1", "-cmd", "@SUM(A1:A2)")
    unsafe_urls = ("javascript:alert(1)", "file:///etc/passwd", "data:text/html,x")
    filenames = ("../../lead.xlsx", "/var/tmp/a.xlsx", "..\\..\\lead.xlsx")
    formula_correct = sum(safe_spreadsheet_text(value).startswith("'") for value in formulas)
    url_correct = sum(safe_http_url(value) is None for value in unsafe_urls)
    filename_correct = 0
    for value in filenames:
        cleaned = sanitize_filename(value)
        filename_correct += bool(
            "/" not in cleaned
            and "\\" not in cleaned
            and ".." not in cleaned
            and cleaned.endswith(".xlsx")
        )
    return len(formulas) + len(unsafe_urls) + len(filenames), formula_correct, url_correct, filename_correct


def _ui_excel_consistency(deps, task, bundle) -> tuple[int, int, float, float, float]:
    request = CreateExportRequest(
        task_id=task.task_id,
        task_version=task.version,
        snapshot_id=bundle.snapshot.snapshot_id,
        fields=[
            "rank",
            "enterprise_name",
            "lead_score",
            "verification_status",
            "public_phone",
        ],
    )
    first = deps.export_service.create_export(request)
    second = deps.export_service.create_export(request)
    workbook = load_workbook(first.artifact_path, data_only=False)
    sheet = workbook["潜客清单"]
    correct = total = 0
    keys = request.fields
    for row_number, lead in enumerate(bundle.leads, start=2):
        for column, key in enumerate(keys, start=1):
            total += 1
            correct += _cell_matches(
                sheet.cell(row_number, column).value,
                _cell_value(getattr(lead, key)),
            )
    invalid_cases = (["unknown"], ["rank", "rank"], ["private_contact"])
    rejected = 0
    for fields in invalid_cases:
        try:
            deps.export_service.registry.validate(fields)
        except ValueError:
            rejected += 1
    return (
        total,
        correct,
        float(first.export_id == second.export_id),
        float(workbook.sheetnames == ["潜客清单", "评分说明", "证据摘要", "任务信息", "冲突明细"]),
        rejected / len(invalid_cases),
    )


def run_evaluation() -> DeliveryEvalReport:
    deps, task, bundle = _completed_bundle("phase7-delivery-eval")
    snapshot_count, snapshot_correct = _snapshot_cases(deps, bundle)
    field_count, field_correct = _export_field_cases(bundle)
    evidence_count, evidence_correct = _evidence_cases(bundle)
    consistency_count, consistency_correct, idempotency, parse_success, allowlist = (
        _ui_excel_consistency(deps, task, bundle)
    )
    historical_count, historical_correct = _historical_cases(deps, task, bundle)
    security_count, formula_correct, url_correct, filename_correct = _security_cases()
    return DeliveryEvalReport(
        snapshot_case_count=snapshot_count,
        export_field_case_count=field_count,
        evidence_case_count=evidence_count,
        historical_case_count=historical_count,
        security_case_count=security_count,
        snapshot_version_accuracy=snapshot_correct / snapshot_count,
        ui_excel_critical_field_consistency=consistency_correct / consistency_count,
        export_field_accuracy=field_correct / field_count,
        export_field_allowlist_enforcement=allowlist,
        evidence_traceability=evidence_correct / evidence_count,
        export_idempotency=idempotency,
        workbook_parse_success=parse_success,
        unsafe_hyperlink_rejection=url_correct / 3,
        formula_injection_protection=formula_correct / 4,
        filename_traversal_protection=filename_correct / 3,
        historical_version_accuracy=historical_correct / historical_count,
    )


if __name__ == "__main__":
    print(json.dumps(asdict(run_evaluation()), ensure_ascii=False, indent=2))
