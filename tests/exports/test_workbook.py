from io import BytesIO

from openpyxl import load_workbook

from app.agent.dependencies import build_dependencies
from app.agent.graph import build_main_graph
from app.delivery.models import DeliveryBundle
from app.domain.task import TaskPatch
from app.exports.models import CreateExportRequest
from app.exports.registry import ExportFieldRegistry
from app.exports.workbook import build_workbook


def _completed_dependencies(session_id: str):
    deps = build_dependencies()
    graph = build_main_graph(deps)
    graph.invoke(
        {"session_id": session_id, "incoming_text": "帮我找上海松江3家集团V网企业"},
        config={"configurable": {"thread_id": session_id}},
    )
    return deps


def _expanded_bundle(bundle: DeliveryBundle, count: int) -> DeliveryBundle:
    leads = []
    details = {}
    for index in range(count):
        source = bundle.leads[index % len(bundle.leads)]
        enterprise_id = f"{source.enterprise_id}-{index}"
        lead = source.model_copy(
            deep=True,
            update={
                "rank": index + 1,
                "enterprise_id": enterprise_id,
                "enterprise_name": f"{source.enterprise_name}-{index + 1}",
            },
        )
        detail = bundle.details[source.enterprise_id].model_copy(
            deep=True,
            update={
                "lead": lead,
                "score": bundle.details[source.enterprise_id].score.model_copy(
                    deep=True, update={"enterprise_id": enterprise_id}
                ),
            },
        )
        leads.append(lead)
        details[enterprise_id] = detail
    return bundle.model_copy(
        deep=True,
        update={
            "snapshot": bundle.snapshot.model_copy(update={"result_count": count}),
            "leads": leads,
            "details": details,
        },
    )


def test_workbook_structure_types_and_ui_projection_are_identical():
    deps = _completed_dependencies("workbook")
    task = deps.task_repository.get_active_task("workbook")
    bundle = deps.delivery_query_service.freeze(task.task_id, task.version)
    job = deps.export_service.create_export(
        CreateExportRequest(
            task_id=task.task_id,
            task_version=task.version,
            snapshot_id=bundle.snapshot.snapshot_id,
            fields=["rank", "enterprise_name", "lead_score", "verification_status"],
        )
    )

    workbook = load_workbook(job.artifact_path)
    assert workbook.sheetnames == ["潜客清单", "评分说明", "证据摘要", "任务信息", "冲突明细"]
    leads = workbook["潜客清单"]
    assert leads.freeze_panes == "A2"
    assert leads.auto_filter.ref == leads.dimensions
    assert leads.max_row == job.row_count + 1
    assert isinstance(leads["A2"].value, int)
    assert isinstance(leads["C2"].value, float)
    assert [leads.cell(2, column).value for column in range(1, 5)] == [
        bundle.leads[0].rank,
        bundle.leads[0].enterprise_name,
        bundle.leads[0].lead_score,
        bundle.leads[0].verification_status,
    ]
    task_sheet = workbook["任务信息"]
    metadata = {task_sheet.cell(row, 1).value: task_sheet.cell(row, 2).value for row in range(2, task_sheet.max_row + 1)}
    assert metadata["task_version"] == task.version
    assert metadata["delivery_snapshot_id"] == bundle.snapshot.snapshot_id


def test_export_is_idempotent_and_bound_to_frozen_historical_snapshot():
    deps = _completed_dependencies("export-idempotent")
    task = deps.task_repository.get_active_task("export-idempotent")
    bundle = deps.delivery_query_service.freeze(task.task_id, task.version)
    request = CreateExportRequest(
        task_id=task.task_id,
        task_version=task.version,
        snapshot_id=bundle.snapshot.snapshot_id,
        fields=["rank", "enterprise_name", "lead_score", "verification_status"],
    )
    first = deps.export_service.create_export(request)
    deps.task_service.apply_patch(task.task_id, TaskPatch(target_count=2))
    second = deps.export_service.create_export(request)

    assert first.export_id == second.export_id
    workbook = load_workbook(first.artifact_path)
    metadata = {workbook["任务信息"].cell(row, 1).value: workbook["任务信息"].cell(row, 2).value for row in range(2, workbook["任务信息"].max_row + 1)}
    assert metadata["task_version"] == bundle.snapshot.task_version
    assert workbook["潜客清单"].max_row == len(bundle.leads) + 1


def test_formula_injection_is_written_as_text_and_unsafe_url_is_not_linked():
    deps = _completed_dependencies("workbook-security")
    task = deps.task_repository.get_active_task("workbook-security")
    execution = deps.execution_snapshot_repository.current(task.task_id)
    lead_set = deps.lead_score_repository.get_lead_set(execution.lead_score_set_id)
    enterprise_id = lead_set.lead_ids[0]
    enterprise = deps.enterprise_repository.enterprises[enterprise_id]
    deps.enterprise_repository.enterprises[enterprise_id] = enterprise.model_copy(
        update={
            "canonical_name": '=HYPERLINK("https://evil.example")',
            "primary_website": "javascript:alert(1)",
        }
    )
    for profile_id, profile in list(deps.evidence_repository.profiles.items()):
        if profile.enterprise_id != enterprise_id or not profile.field("website"):
            continue
        fields = dict(profile.fields)
        fields["website"] = fields["website"].model_copy(
            update={"primary_value": "javascript:alert(1)"}
        )
        deps.evidence_repository.profiles[profile_id] = profile.model_copy(
            update={"fields": fields}
        )
    bundle = deps.delivery_query_service.freeze(task.task_id, task.version, force_new=True)
    job = deps.export_service.create_export(
        CreateExportRequest(
            task_id=task.task_id,
            snapshot_id=bundle.snapshot.snapshot_id,
            fields=["enterprise_name", "website", "verification_status"],
        )
    )
    workbook = load_workbook(job.artifact_path)
    sheet = workbook["潜客清单"]
    assert sheet["A2"].data_type == "s"
    assert sheet["A2"].value.startswith("'=")
    assert sheet["B2"].hyperlink is None


def test_workbook_with_1000_rows_is_generated_and_parseable():
    deps = _completed_dependencies("workbook-1000")
    task = deps.task_repository.get_active_task("workbook-1000")
    bundle = deps.delivery_query_service.freeze(task.task_id, task.version)
    expanded = _expanded_bundle(bundle, 1000)
    fields = ExportFieldRegistry().validate(
        ["rank", "enterprise_name", "industry", "lead_score", "verification_status"]
    )

    content = build_workbook(
        expanded,
        fields,
        target_count=1000,
        export_id="workbook-1000",
        exported_at=bundle.snapshot.created_at,
        include_task_summary=True,
        include_score_breakdown=True,
        include_evidence_summary=True,
        include_conflicts=True,
    )
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False)

    assert workbook.sheetnames == ["潜客清单", "评分说明", "证据摘要", "任务信息", "冲突明细"]
    assert workbook["潜客清单"].max_row == 1001
    assert workbook["评分说明"].max_row == 1001
    assert workbook["证据摘要"].max_row > 1000
    assert workbook["潜客清单"]["A1001"].value == 1000
